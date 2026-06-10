from __future__ import annotations

from pathlib import Path
from typing import Any

from ..models import DocumentMetadata, ParsedDocument, ParsedTable
from .base import BaseParser


class XlsxParser(BaseParser):
    """
    Parser backend для книг Excel (.xlsx / .xlsm).

    Каждый непустой лист книги превращается в одну ParsedTable
    (markdown + raw_data). Очень большой лист режется на окна по строкам,
    чтобы один чанк не получался гигантским.

    Логика универсальна — никаких привязок к конкретному формату отчёта.
    Просто читаем сетку ячеек, убираем полностью пустые строки и столбцы.

    openpyxl импортируется лениво внутри parse(), чтобы отсутствие пакета
    не ломало запуск всего пайплайна: обработка PDF от openpyxl не зависит.
    """

    parser_name: str = "xlsx"
    parser_version: str = "0.1.0"
    supported_extensions: set[str] = {".xlsx", ".xlsm"}

    # Лист крупнее этого числа строк (после очистки) режется на окна.
    max_rows_per_table: int = 200

    def __init__(self, max_rows_per_table: int | None = None) -> None:
        if max_rows_per_table is not None:
            if max_rows_per_table < 2:
                raise ValueError("max_rows_per_table must be >= 2")
            self.max_rows_per_table = max_rows_per_table

    # ============================================================
    # PUBLIC
    # ============================================================

    def parse(
        self,
        path: Path,
        metadata: DocumentMetadata,
    ) -> ParsedDocument:
        metadata.parser_name = self.parser_name
        metadata.parser_version = self.parser_version

        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError(
                "XlsxParser требует пакет 'openpyxl'. "
                "Установите его: pip install openpyxl"
            ) from exc

        tables: list[ParsedTable] = []
        table_order = 0

        # read_only — экономно по памяти; data_only — берём значения, а не формулы.
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)

        try:
            for sheet_index, worksheet in enumerate(workbook.worksheets):
                grid = self._read_sheet_grid(worksheet)
                if not grid:
                    continue

                for window in self._split_into_windows(grid):
                    table = self._make_table(
                        rows=window,
                        table_order=table_order,
                        sheet_name=worksheet.title,
                        sheet_index=sheet_index,
                        metadata=metadata,
                    )
                    if table is not None:
                        tables.append(table)
                        table_order += 1
        finally:
            workbook.close()

        return ParsedDocument(
            metadata=metadata,
            blocks=[],
            tables=tables,
            visuals=[],
        )

    # ============================================================
    # SHEET -> CLEAN GRID
    # ============================================================

    def _read_sheet_grid(self, worksheet) -> list[list[str]]:
        """
        Читает лист в прямоугольную сетку строк одинаковой ширины.

        Убирает полностью пустые строки и полностью пустые столбцы
        (выгрузки часто содержат десятки пустых столбцов справа).
        """

        raw_rows: list[tuple[Any, ...]] = list(worksheet.iter_rows(values_only=True))
        if not raw_rows:
            return []

        width = max((len(row) for row in raw_rows), default=0)
        if width == 0:
            return []

        # Приводим ячейки к строкам, дополняем строки до общей ширины.
        rows: list[list[str]] = []
        for raw in raw_rows:
            cells = [self._clean_cell(c) for c in raw]
            if len(cells) < width:
                cells += [""] * (width - len(cells))
            rows.append(cells)

        # Какие столбцы непустые хотя бы в одной строке.
        keep_cols = [
            col
            for col in range(width)
            if any(rows[r][col] for r in range(len(rows)))
        ]
        if not keep_cols:
            return []

        # Оставляем только непустые столбцы и непустые строки.
        cleaned: list[list[str]] = []
        for row in rows:
            new_row = [row[col] for col in keep_cols]
            if any(cell for cell in new_row):
                cleaned.append(new_row)

        return cleaned

    @staticmethod
    def _clean_cell(value: Any) -> str:
        """Приводит значение ячейки к аккуратной строке."""

        if value is None:
            return ""

        # Целые float-ы (7976290.0) приводим к int, чтобы не было хвоста .0.
        if isinstance(value, float) and value.is_integer():
            value = int(value)

        text = str(value)
        text = text.replace("\r", " ").replace("\n", " ")
        text = text.replace("|", "/")  # символ | ломает markdown-таблицу
        text = " ".join(text.split())
        return text.strip()

    # ============================================================
    # WINDOWING
    # ============================================================

    def _split_into_windows(
        self,
        grid: list[list[str]],
    ) -> list[list[list[str]]]:
        """
        Небольшой лист отдаём одним куском. Очень большой режем на окна
        по строкам; первая строка (шапка) повторяется в каждом окне.
        """

        if len(grid) <= self.max_rows_per_table:
            return [grid]

        header = grid[0]
        body = grid[1:]
        step = max(1, self.max_rows_per_table - 1)  # минус строка шапки

        windows: list[list[list[str]]] = []
        for start in range(0, len(body), step):
            chunk = body[start:start + step]
            windows.append([header] + chunk)
        return windows

    # ============================================================
    # GRID -> ParsedTable
    # ============================================================

    def _make_table(
        self,
        rows: list[list[str]],
        table_order: int,
        sheet_name: str,
        sheet_index: int,
        metadata: DocumentMetadata,
    ) -> ParsedTable | None:
        if not rows or not rows[0]:
            return None

        markdown = self._to_markdown(rows)
        if not markdown.strip():
            return None

        return ParsedTable(
            markdown_table=markdown,
            raw_data=[list(row) for row in rows],
            table_order=table_order,
            source_file=metadata.source_file,
            source_type=metadata.source_type,
            # У Excel нет страниц — как локатор используем индекс листа (1-based).
            page_number=sheet_index + 1,
            table_title=sheet_name,
            language=metadata.language,
            parser_name=self.parser_name,
            parser_version=self.parser_version,
            metadata={
                "project": metadata.project,
                "market": metadata.market,
                "product": metadata.product,
                "parser_backend": self.parser_name,
                "parser_name": self.parser_name,
                "parser_version": self.parser_version,
                "table_format": "markdown",
                "from_xlsx_sheet": True,
                "sheet_name": sheet_name,
                "sheet_index": sheet_index,
                # У Excel-книги нет «разделов», но имя листа — лучший
                # доступный контекст для retrieval.
                "section_path": sheet_name,
            },
        )

    @staticmethod
    def _to_markdown(rows: list[list[str]]) -> str:
        """Собирает валидную markdown-таблицу из сетки строк равной ширины."""

        if not rows:
            return ""

        width = len(rows[0])
        if width == 0:
            return ""

        def fmt(row: list[str]) -> str:
            cells = list(row) + [""] * (width - len(row))
            cells = cells[:width]
            return "| " + " | ".join(cell or " " for cell in cells) + " |"

        lines = [
            fmt(rows[0]),
            "| " + " | ".join(["---"] * width) + " |",
        ]
        for row in rows[1:]:
            lines.append(fmt(row))

        return "\n".join(lines)
